"""Import and export utilities for price lists."""

from __future__ import annotations

import csv
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Sequence


class MissingRequiredColumnsError(ValueError):
    """Raised when required columns cannot be resolved during import."""

    def __init__(
        self,
        missing: Sequence[str],
        path: Path,
        headers: Sequence[str],
        resolved: dict[str, str],
    ) -> None:
        self.missing = tuple(missing)
        self.path = Path(path)
        self.headers = tuple(headers)
        self.resolved = dict(resolved)
        message = (
            "Missing required columns {} in file '{}'. You can provide a column_mapping "
            "or extend column_aliases to match supplier headers.".format(
                list(missing), path
            )
        )
        super().__init__(message)

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]

from .models import PriceList, Product

SUPPORTED_IMPORT_FORMATS = {".csv", ".json", ".xlsx"}
SUPPORTED_EXPORT_FORMATS = {".csv", ".json", ".xlsx"}


class PriceListImporter:
    """Load price lists from supported file formats."""

    def __init__(
        self,
        default_currency: str = "USD",
        column_aliases: dict[str, Iterable[str]] | None = None,
    ) -> None:
        self.default_currency = default_currency
        self.required_fields = {"sku", "name", "price"}

        default_aliases: dict[str, set[str]] = {
            "sku": {
                "sku",
                "артикул",
                "article",
                "код товару",
                "код",
                "товарний код",
            },
            "name": {
                "name",
                "назва",
                "назва позиції",
                "найменування",
                "товар",
                "product",
            },
            "price": {
                "price",
                "ціна",
                "вартість",
                "cost",
                "цена",
            },
            "currency": {"currency", "валюта"},
            "description": {"description", "опис", "описання"},
            "tags": {"tags", "теги", "мітки"},
        }

        if column_aliases:
            for field, aliases in column_aliases.items():
                normalized = {str(alias).strip().lower() for alias in aliases}
                default_aliases.setdefault(field, set()).update(normalized)

        for field in list(default_aliases):
            default_aliases[field].add(field)

        self.column_aliases = {
            field: tuple(sorted({alias.strip().lower() for alias in aliases if alias}))
            for field, aliases in default_aliases.items()
        }

        self.optional_fields = tuple(
            sorted(field for field in self.column_aliases if field not in self.required_fields)
        )

    def load(
        self,
        path: str | Path,
        supplier: str | None = None,
        column_mapping: dict[str, str | Sequence[str]] | None = None,
    ) -> PriceList:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_IMPORT_FORMATS:
            raise ValueError(f"Unsupported import format '{suffix}'.")

        if suffix == ".csv":
            products = list(self._load_csv(path, supplier, column_mapping))
        elif suffix == ".json":
            products = list(self._load_json(path, supplier, column_mapping))
        else:
            products = list(self._load_excel(path, supplier, column_mapping))

        supplier_name = supplier or path.stem
        price_list = PriceList(supplier=supplier_name, products=products)
        price_list.sort_products()
        return price_list

    def detect_headers(self, path: str | Path) -> list[str]:
        """Return a list of header names found in the supplied file."""

        headers, rows = self.peek(path, limit=0)
        if headers:
            return headers
        if rows:
            # When rows are provided we can derive headers from the first row keys.
            first_row = rows[0]
            return list(first_row.keys())

        path = Path(path)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            dialect, delimiter = self._detect_csv_format(path)
            with path.open("r", encoding="utf-8-sig", newline="") as fp:
                reader = (
                    csv.reader(fp, dialect=dialect)
                    if dialect is not None
                    else csv.reader(fp, delimiter=delimiter)
                )
                try:
                    headers = next(reader)
                except StopIteration:
                    return []
            return [str(header).strip() for header in headers if str(header).strip()]

        if suffix == ".json":
            with path.open("r", encoding="utf-8") as fp:
                payload = json.load(fp)
            products_data = payload.get("products") if isinstance(payload, dict) else None
            if not isinstance(products_data, list):
                return []
            headers: list[str] = []
            for item in products_data:
                if not isinstance(item, dict):
                    continue
                for key in item.keys():
                    key_str = str(key)
                    if key_str not in headers:
                        headers.append(key_str)
            return headers

        if suffix == ".xlsx":
            if load_workbook is None:
                raise ValueError("Імпорт Excel недоступний. Встановіть залежність 'openpyxl'.")
            workbook = load_workbook(path, data_only=True)
            sheet = workbook.active
            rows_iter = sheet.iter_rows(min_row=1, values_only=True)
            try:
                headers_row = next(rows_iter)
            except StopIteration:
                return []
            headers = [
                str(value).strip() if value is not None else ""
                for value in headers_row
            ]
            return [header for header in headers if header]

        return []

    def peek(
        self, path: str | Path, *, limit: int = 10
    ) -> tuple[list[str], list[dict[str, str]]]:
        """Return headers and a preview of rows from the provided file."""

        path = Path(path)
        suffix = path.suffix.lower()
        headers: list[str] = []
        rows: list[dict[str, str]] = []

        if suffix == ".csv":
            dialect, delimiter = self._detect_csv_format(path)
            with path.open("r", encoding="utf-8-sig", newline="") as fp:
                reader = (
                    csv.DictReader(fp, dialect=dialect)
                    if dialect is not None
                    else csv.DictReader(fp, delimiter=delimiter)
                )
                headers = [
                    str(header).strip()
                    for header in (reader.fieldnames or [])
                    if header is not None and str(header).strip()
                ]
                if limit == 0:
                    return headers, rows
                for index, row in enumerate(reader):
                    if limit and index >= limit:
                        break
                    rows.append({header: str(row.get(header, "")) for header in headers})
            return headers, rows

        if suffix == ".json":
            with path.open("r", encoding="utf-8") as fp:
                payload = json.load(fp)

            products_data = payload.get("products") if isinstance(payload, dict) else None
            if isinstance(products_data, list):
                for item in products_data:
                    if not isinstance(item, dict):
                        continue
                    for key in item.keys():
                        key_str = str(key)
                        if key_str not in headers:
                            headers.append(key_str)
                if limit == 0:
                    return headers, rows
                for item in products_data[: limit or None]:
                    if not isinstance(item, dict):
                        continue
                    rows.append({header: str(item.get(header, "")) for header in headers})
            return headers, rows

        if suffix == ".xlsx":
            if load_workbook is None:
                raise ValueError("Імпорт Excel недоступний. Встановіть залежність 'openpyxl'.")

            workbook = load_workbook(path, data_only=True)
            sheet = workbook.active
            rows_iter = sheet.iter_rows(min_row=1, values_only=True)

            try:
                headers_row = next(rows_iter)
            except StopIteration:
                return headers, rows

            headers = [
                str(value).strip() if value is not None else ""
                for value in headers_row
            ]
            headers = [header for header in headers if header]

            if limit == 0:
                return headers, rows

            for index, values in enumerate(rows_iter):
                if limit and index >= limit:
                    break
                row_dict = {}
                for idx, header in enumerate(headers):
                    if idx < len(values):
                        value = values[idx]
                    else:
                        value = ""
                    row_dict[header] = "" if value is None else str(value)
                rows.append(row_dict)
            return headers, rows

        raise ValueError(f"Unsupported import format '{suffix}'.")

    def suggest_mapping(
        self,
        headers: Sequence[str],
        column_mapping: dict[str, str | Sequence[str]] | None = None,
    ) -> dict[str, str]:
        """Suggest column mapping for given headers without enforcing required fields."""

        normalized_headers = {
            str(header).strip().lower(): str(header)
            for header in headers
            if header is not None and str(header).strip()
        }

        resolved: dict[str, str] = {}

        def iter_override(field: str) -> Iterable[str]:
            if not column_mapping or field not in column_mapping:
                return []
            value = column_mapping[field]
            if isinstance(value, str):
                return [value]
            return [str(candidate) for candidate in value]

        for field, aliases in self.column_aliases.items():
            override_candidates = [
                str(candidate).strip().lower()
                for candidate in iter_override(field)
                if str(candidate).strip()
            ]
            candidates = override_candidates + [
                alias for alias in aliases if alias not in override_candidates
            ]

            for candidate in candidates:
                candidate_key = candidate.strip().lower()
                if not candidate_key:
                    continue
                if candidate_key in normalized_headers:
                    resolved[field] = normalized_headers[candidate_key]
                    break

        return resolved

    def _load_csv(
        self,
        path: Path,
        supplier: str | None,
        column_mapping: dict[str, str | Sequence[str]] | None,
    ) -> Iterable[Product]:
        dialect, delimiter = self._detect_csv_format(path)
        with path.open("r", encoding="utf-8-sig", newline="") as fp:
            reader = (
                csv.DictReader(fp, dialect=dialect)
                if dialect is not None
                else csv.DictReader(fp, delimiter=delimiter)
            )
            headers = reader.fieldnames or []
            mapping = self._prepare_column_mapping(headers, column_mapping, path)
            used_columns = set(mapping.values())

            for row in reader:
                sku = self._extract_cell(row, mapping["sku"]).strip()
                name = self._extract_cell(row, mapping["name"]).strip()
                price_raw = self._extract_cell(row, mapping["price"])
                price_str = str(price_raw).strip() if price_raw is not None else ""

                currency_column = mapping.get("currency")
                currency_value = (
                    self._extract_cell(row, currency_column) if currency_column else ""
                )
                currency = currency_value.strip() or self.default_currency

                description_column = mapping.get("description")
                description_value = (
                    self._extract_cell(row, description_column) if description_column else ""
                )
                description = description_value.strip() or None

                tags_column = mapping.get("tags")
                tags_raw = self._extract_cell(row, tags_column) if tags_column else ""

                if not sku or not name or not price_str:
                    continue

                price = self._parse_price(price_raw, sku)

                tags = {
                    tag.strip().lower()
                    for tag in str(tags_raw).split(";")
                    if str(tag).strip()
                }

                extra = {
                    key: value
                    for key, value in row.items()
                    if key not in used_columns and value not in {None, ""}
                }

                yield Product(
                    sku=sku,
                    name=name,
                    price=price,
                    currency=currency,
                    description=description,
                    supplier=supplier,
                    tags=tags,
                    extra=extra,
                )

    def _load_json(
        self,
        path: Path,
        supplier: str | None,
        column_mapping: dict[str, str | Sequence[str]] | None,
    ) -> Iterable[Product]:
        with path.open("r", encoding="utf-8") as fp:
            payload = json.load(fp)

        products_data = payload.get("products")
        if not isinstance(products_data, list):
            raise ValueError("JSON payload must contain a 'products' list.")

        header_index: dict[str, str] = {}
        for item in products_data:
            if isinstance(item, dict):
                for key in item.keys():
                    key_str = str(key)
                    normalized = key_str.strip().lower()
                    header_index.setdefault(normalized, key_str)

        if not header_index and not products_data:
            return

        mapping = self._prepare_column_mapping(
            list(header_index.values()), column_mapping, path
        )
        used_columns = set(mapping.values())

        for item in products_data:
            if not isinstance(item, dict):
                continue

            normalized_row = {
                str(key).strip().lower(): value for key, value in item.items()
            }

            def _get(field: str) -> object | None:
                column = mapping.get(field)
                if not column:
                    return None
                if column in item and item[column] is not None:
                    return item[column]
                column_lower = column.strip().lower()
                if column_lower in normalized_row and normalized_row[column_lower] is not None:
                    return normalized_row[column_lower]
                return None

            sku = str(_get("sku") or "").strip()
            name = str(_get("name") or "").strip()
            price_raw = _get("price")
            price_str = str(price_raw).strip() if price_raw is not None else ""

            if not sku or not name or not price_str:
                continue

            price = self._parse_price(price_raw, sku)

            currency_value = _get("currency")
            currency = str(currency_value).strip() if currency_value not in {None, ""} else ""
            currency = currency or self.default_currency

            description_value = _get("description")
            if isinstance(description_value, str):
                description = description_value.strip() or None
            elif description_value is None:
                description = None
            else:
                description = str(description_value)

            tags_value = _get("tags")
            tags: set[str]
            if isinstance(tags_value, (list, tuple, set)):
                tags = {
                    str(tag).strip().lower()
                    for tag in tags_value
                    if str(tag).strip()
                }
            else:
                tags_raw = str(tags_value).strip() if tags_value is not None else ""
                tags = {
                    tag.strip().lower()
                    for tag in tags_raw.split(";")
                    if tag.strip()
                }

            if not tags and isinstance(item.get("tags"), list):
                tags = {
                    str(tag).strip().lower()
                    for tag in item.get("tags", [])
                    if str(tag).strip()
                }

            extra = {
                key: value
                for key, value in item.items()
                if key not in used_columns and key not in {"extra"}
            }
            extra_payload = item.get("extra")
            if isinstance(extra_payload, dict):
                extra.update(extra_payload)

            yield Product(
                sku=sku,
                name=name,
                price=price,
                currency=currency,
                description=description,
                supplier=supplier,
                tags=tags,
                extra=extra,
            )

    def _load_excel(
        self,
        path: Path,
        supplier: str | None,
        column_mapping: dict[str, str | Sequence[str]] | None,
    ) -> Iterable[Product]:
        if load_workbook is None:
            raise ValueError(
                "Імпорт Excel недоступний. Встановіть залежність 'openpyxl'."
            )

        workbook = load_workbook(path, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(min_row=1, values_only=True)

        try:
            headers_row = next(rows)
        except StopIteration:
            return

        headers = [str(value).strip() if value is not None else "" for value in headers_row]
        mapping = self._prepare_column_mapping(headers, column_mapping, path)
        used_columns = set(mapping.values())

        for row_values in rows:
            if not any(row_values):
                continue

            row = {
                headers[index]: row_values[index]
                for index in range(len(headers))
                if headers[index]
            }
            sku = self._extract_cell(row, mapping["sku"]).strip()
            name = self._extract_cell(row, mapping["name"]).strip()
            price_value = self._extract_cell(row, mapping["price"])
            price_str = str(price_value).strip()

            if not sku or not name or not price_str:
                continue

            currency_column = mapping.get("currency")
            currency_value = self._extract_cell(row, currency_column) if currency_column else ""
            currency = currency_value.strip() or self.default_currency

            description_column = mapping.get("description")
            description_raw = (
                self._extract_cell(row, description_column) if description_column else ""
            )
            description = description_raw.strip() or None

            tags_column = mapping.get("tags")
            tags_raw = self._extract_cell(row, tags_column) if tags_column else ""

            price = self._parse_price(price_value, sku)

            tags = {
                tag.strip().lower()
                for tag in str(tags_raw).split(";")
                if str(tag).strip()
            }

            extra = {
                key: value
                for key, value in row.items()
                if key not in used_columns and value not in {None, ""}
            }

            yield Product(
                sku=sku,
                name=name,
                price=price,
                currency=currency,
                description=description,
                supplier=supplier,
                tags=tags,
                extra={k: str(v) for k, v in extra.items()},
            )

    def _prepare_column_mapping(
        self,
        headers: Sequence[str],
        column_mapping: dict[str, str | Sequence[str]] | None,
        path: Path,
    ) -> dict[str, str]:
        normalized_headers = {
            str(header).strip().lower(): str(header)
            for header in headers
            if header is not None and str(header).strip()
        }

        def iter_override(field: str) -> Iterable[str]:
            if not column_mapping or field not in column_mapping:
                return []
            value = column_mapping[field]
            if isinstance(value, str):
                return [value]
            return [str(candidate) for candidate in value]

        resolved: dict[str, str] = {}
        for field, aliases in self.column_aliases.items():
            override_candidates = [
                str(candidate).strip().lower()
                for candidate in iter_override(field)
                if str(candidate).strip()
            ]
            candidates = override_candidates + [
                alias for alias in aliases if alias not in override_candidates
            ]

            for candidate in candidates:
                candidate_key = candidate.strip().lower()
                if not candidate_key:
                    continue
                if candidate_key in normalized_headers:
                    resolved[field] = normalized_headers[candidate_key]
                    break

        missing = sorted(self.required_fields - set(resolved))
        if missing:
            available_headers = [
                str(header)
                for header in headers
                if header is not None and str(header).strip()
            ]
            raise MissingRequiredColumnsError(
                missing=missing,
                path=path,
                headers=available_headers,
                resolved=resolved,
            )

        return resolved

    def _detect_csv_format(self, path: Path) -> tuple[csv.Dialect | None, str]:
        """Detect CSV dialect and delimiter, falling back to sensible defaults."""

        sample = ""
        with path.open("r", encoding="utf-8-sig", newline="") as fp:
            sample = fp.read(4096)

        dialect: csv.Dialect | None = None
        if sample:
            try:
                dialect = csv.Sniffer().sniff(sample)
            except csv.Error:
                dialect = None

        if dialect is not None:
            delimiter = getattr(dialect, "delimiter", ",") or ","
            return dialect, delimiter

        delimiter = self._guess_delimiter(sample)
        return None, delimiter

    def _parse_price(self, price_raw: object, sku: str) -> float:
        """Parse numeric price values from a wide range of string formats."""

        if isinstance(price_raw, (int, float, Decimal)):
            return float(price_raw)

        if price_raw is None:
            raise ValueError(f"Invalid price '{price_raw}' for SKU '{sku}'.")

        price_str = str(price_raw).strip()
        if not price_str:
            raise ValueError(f"Invalid price '{price_raw}' for SKU '{sku}'.")

        normalized = price_str.replace("\u00a0", "").replace(" ", "")
        normalized = normalized.replace("−", "-")
        normalized = re.sub(r"[^0-9,.-]", "", normalized)

        if not normalized:
            raise ValueError(f"Invalid price '{price_raw}' for SKU '{sku}'.")

        sign = ""
        if normalized.startswith("-"):
            sign = "-"
            normalized = normalized[1:]
        normalized = normalized.replace("-", "")
        normalized = f"{sign}{normalized}"

        if not any(char.isdigit() for char in normalized):
            raise ValueError(f"Invalid price '{price_raw}' for SKU '{sku}'.")

        if "," in normalized and "." in normalized:
            if normalized.rfind(",") > normalized.rfind("."):
                normalized = normalized.replace(".", "")
                normalized = normalized.replace(",", ".")
            else:
                normalized = normalized.replace(",", "")
        elif "," in normalized:
            if normalized.count(",") > 1:
                normalized = normalized.replace(",", "")
            else:
                normalized = normalized.replace(",", ".")
        elif normalized.count(".") > 1:
            normalized = normalized.replace(".", "")

        try:
            return float(Decimal(normalized))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError(f"Invalid price '{price_raw}' for SKU '{sku}'.") from exc

    @staticmethod
    def _guess_delimiter(sample: str) -> str:
        candidates = [",", ";", "\t", "|"]
        counts = {candidate: sample.count(candidate) for candidate in candidates}
        best_candidate = max(candidates, key=lambda candidate: counts[candidate])
        if counts[best_candidate] == 0:
            return ","
        return best_candidate

    @staticmethod
    def _extract_cell(row: dict[str, object], column: str | None) -> str:
        if not column:
            return ""
        if column in row and row[column] is not None:
            return str(row[column])
        column_lower = column.strip().lower()
        for key, value in row.items():
            if str(key).strip().lower() == column_lower and value is not None:
                return str(value)
        return ""


class PriceListExporter:
    """Export price list or product collections into supported formats."""

    def export(self, products: Sequence[Product], path: str | Path) -> None:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXPORT_FORMATS:
            raise ValueError(f"Unsupported export format '{suffix}'.")

        if suffix == ".csv":
            self._export_csv(products, path)
        elif suffix == ".json":
            self._export_json(products, path)
        else:
            self._export_excel(products, path)

    def _export_csv(self, products: Sequence[Product], path: Path) -> None:
        fieldnames = ["sku", "name", "price", "currency", "description", "tags", "supplier"]
        extra_fields = sorted({key for product in products for key in product.extra})
        fieldnames.extend(extra_fields)

        with path.open("w", encoding="utf-8", newline="") as fp:
            writer = csv.DictWriter(fp, fieldnames=fieldnames)
            writer.writeheader()
            for product in products:
                row = {
                    "sku": product.sku,
                    "name": product.name,
                    "price": product.price,
                    "currency": product.currency,
                    "description": product.description or "",
                    "tags": ";".join(sorted(product.tags)),
                    "supplier": product.supplier or "",
                }
                for key in extra_fields:
                    row[key] = product.extra.get(key, "")
                writer.writerow(row)

    def _export_json(self, products: Sequence[Product], path: Path) -> None:
        payload = {
            "products": [
                {
                    "sku": product.sku,
                    "name": product.name,
                    "price": product.price,
                    "currency": product.currency,
                    "description": product.description,
                    "tags": sorted(product.tags),
                    "supplier": product.supplier,
                    "extra": product.extra,
                }
                for product in products
            ]
        }
        with path.open("w", encoding="utf-8") as fp:
            json.dump(payload, fp, indent=2, ensure_ascii=False)

    def _export_excel(self, products: Sequence[Product], path: Path) -> None:
        if Workbook is None:
            raise ValueError(
                "Експорт у Excel недоступний. Встановіть залежність 'openpyxl'."
            )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Products"

        headers = ["sku", "name", "price", "currency", "description", "tags", "supplier"]
        extra_fields = sorted({key for product in products for key in product.extra})
        headers.extend(extra_fields)

        sheet.append(headers)

        for product in products:
            row = [
                product.sku,
                product.name,
                product.price,
                product.currency,
                product.description or "",
                ";".join(sorted(product.tags)),
                product.supplier or "",
            ]
            row.extend(product.extra.get(key, "") for key in extra_fields)
            sheet.append(row)

        workbook.save(path)
