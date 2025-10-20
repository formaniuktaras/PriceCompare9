"""Utilities for managing hierarchical model templates used for tagging."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, MutableMapping, Optional, Sequence

try:  # pragma: no cover - optional dependency
    from openpyxl import Workbook, load_workbook
except Exception:  # pragma: no cover - optional dependency
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]


_EXCEL_HEADERS = (
    "Категорія",
    "Теги категорії",
    "Бренд",
    "Теги бренду",
    "Модель",
    "Теги моделі",
    "Патерн",
)


def _split_tags(value: object) -> List[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        parts = [str(item).strip() for item in value if str(item).strip()]
    else:
        raw = str(value)
        parts = [part.strip() for part in raw.replace("\n", ";").replace(",", ";").split(";")]
        parts = [part for part in parts if part]
    return parts


def _normalize_tags(tags: Iterable[str]) -> List[str]:
    """Return a list with duplicate tags removed while preserving order."""

    seen_lower: set[str] = set()
    ordered: List[str] = []
    for tag in tags:
        if tag is None:
            continue
        cleaned = str(tag).strip()
        lowered = cleaned.lower()
        if not cleaned or lowered in seen_lower:
            continue
        seen_lower.add(lowered)
        ordered.append(cleaned)
    return ordered


def generate_model_pattern(brand: str, model: str) -> str:
    """Generate a flexible regex for matching brand/model combinations.

    The previous implementation escaped the full name verbatim which meant the
    resulting expression only matched the exact spacing present in the
    template. Real-world supplier names often omit spaces or replace them with
    punctuation (for example, ``Samsung GalaxyS23 FE`` or ``Galaxy-S23``).

    To make matching more robust we split the combined brand/model value into
    tokens and allow any number of non-word separators (including the absence
    of one) between them. We still keep the leading/trailing word boundaries so
    that partial matches (like ``S23`` inside ``S230``) are avoided.
    """

    parts = [str(brand or "").strip(), str(model or "").strip()]
    full_name = " ".join(part for part in parts if part)
    if not full_name:
        return r""

    # Allow for missing or substituted separators between tokens (spaces,
    # dashes, underscores, etc.). ``\W`` does not include underscore, so we add
    # it explicitly to cover identifiers that use it as a separator.
    tokens = [token for token in re.split(r"\s+", full_name) if token]
    if not tokens:
        return r""

    escaped_tokens = [re.escape(token) for token in tokens]
    flexible_separator = r"[\W_]*"
    pattern_body = flexible_separator.join(escaped_tokens)
    return rf"(?<![A-Za-z0-9]){pattern_body}(?![A-Za-z0-9])"


@dataclass
class ModelRecord:
    """Represents a single model entry inside the template hierarchy."""

    category: str
    brand: str
    name: str
    pattern: str
    tags: Sequence[str]

    def to_payload(self) -> Dict[str, object]:
        return {
            "name": self.name,
            "pattern": self.pattern,
            "tags": list(self.tags),
        }


class ModelTemplatesEditor:
    """Load, manipulate and persist hierarchical model template data."""

    def __init__(self, filename: str | Path = "models.json") -> None:
        self.path = Path(filename)
        self.data: Dict[str, object] = {"categories": []}
        self._regex_cache: Dict[str, re.Pattern[str]] = {}
        self.load_data()

    # ------------------------------------------------------------------
    # Persistence helpers
    # ------------------------------------------------------------------
    def load_data(self) -> None:
        """Load template data from disk, creating defaults if necessary."""

        if self.path.exists():
            try:
                with self.path.open("r", encoding="utf-8") as fp:
                    payload = json.load(fp)
            except (OSError, json.JSONDecodeError):
                payload = self.default_payload()
        else:
            payload = self.default_payload()

        if not isinstance(payload, MutableMapping):
            payload = {"categories": []}

        categories = payload.get("categories")
        if not isinstance(categories, list):
            categories = []

        normalized = self._normalize_structure(categories)
        self.data = {"categories": normalized}
        existing_categories = payload.get("categories") if isinstance(payload, dict) else None
        if not self.path.exists() or existing_categories != normalized:
            self.save_data()

    def save_data(self) -> None:
        """Persist current data to disk in UTF-8 encoded JSON format."""

        self.path.parent.mkdir(parents=True, exist_ok=True)
        payload = {"categories": self.data.get("categories", [])}
        with self.path.open("w", encoding="utf-8") as fp:
            json.dump(payload, fp, ensure_ascii=False, indent=2)
            fp.write("\n")
        self._regex_cache.clear()

    # ------------------------------------------------------------------
    # Default payload
    # ------------------------------------------------------------------
    @staticmethod
    def default_payload() -> Dict[str, object]:
        """Return a default payload used when no data is stored yet."""

        smartphones = {
            "name": "Смартфони",
            "tags": ["Смартфони", "Смартфон"],
            "brands": [
                {
                    "name": "Apple",
                    "tags": ["Apple"],
                    "models": [
                        {"name": "iPhone 15"},
                        {"name": "iPhone 15 Pro"},
                        {"name": "iPhone 15 Pro Max"},
                    ],
                },
                {
                    "name": "Samsung",
                    "tags": ["Samsung"],
                    "models": [
                        {"name": "Galaxy S23"},
                        {"name": "Galaxy S23 Ultra"},
                    ],
                },
                {
                    "name": "Xiaomi",
                    "tags": ["Xiaomi", "Redmi", "Mi"],
                    "models": [
                        {"name": "Redmi Note 12 4G"},
                        {"name": "Redmi Note 12 Pro"},
                    ],
                },
            ],
        }
        return {"categories": [smartphones]}

    # ------------------------------------------------------------------
    # Normalisation helpers
    # ------------------------------------------------------------------
    def _normalize_structure(self, categories: Iterable[object]) -> List[Dict[str, object]]:
        normalized: List[Dict[str, object]] = []
        for category in categories:
            if not isinstance(category, MutableMapping):
                continue
            name = str(category.get("name") or "").strip()
            if not name:
                continue
            raw_cat_tags = category.get("tags", [name])
            if isinstance(raw_cat_tags, Sequence) and not isinstance(
                raw_cat_tags, (str, bytes)
            ):
                cat_source = list(raw_cat_tags) or [name]
            else:
                cat_source = [raw_cat_tags]
            cat_tags = _normalize_tags(cat_source)
            brands = category.get("brands")
            if not isinstance(brands, list):
                brands = []
            normalized_brands: List[Dict[str, object]] = []
            for brand in brands:
                if not isinstance(brand, MutableMapping):
                    continue
                brand_name = str(brand.get("name") or "").strip()
                if not brand_name:
                    continue
                raw_brand_tags = brand.get("tags", [brand_name])
                if isinstance(raw_brand_tags, Sequence) and not isinstance(
                    raw_brand_tags, (str, bytes)
                ):
                    brand_source = list(raw_brand_tags) or [brand_name]
                else:
                    brand_source = [raw_brand_tags]
                brand_tags = _normalize_tags(brand_source)
                models = brand.get("models")
                if not isinstance(models, list):
                    models = []
                normalized_models: List[Dict[str, object]] = []
                for model in models:
                    if not isinstance(model, MutableMapping):
                        continue
                    model_name = str(model.get("name") or "").strip()
                    if not model_name:
                        continue
                    pattern = str(model.get("pattern") or "").strip()
                    if not pattern:
                        pattern = generate_model_pattern(brand_name, model_name)
                    tags = model.get("tags") or [name, brand_name, model_name]
                    if isinstance(tags, Sequence) and not isinstance(tags, (str, bytes)):
                        extra_tags = list(tags)
                    else:
                        extra_tags = [tags]
                    normalized_tags = _normalize_tags(
                        [name, brand_name, model_name, *extra_tags]
                    )
                    record = {
                        "name": model_name,
                        "pattern": pattern,
                        "tags": normalized_tags,
                    }
                    normalized_models.append(record)
                normalized_models.sort(key=lambda item: item["name"].lower())
                normalized_brands.append(
                    {
                        "name": brand_name,
                        "tags": brand_tags if brand_tags else [brand_name],
                        "models": normalized_models,
                    }
                )
            normalized_brands.sort(key=lambda item: item["name"].lower())
            normalized.append(
                {
                    "name": name,
                    "tags": cat_tags if cat_tags else [name],
                    "brands": normalized_brands,
                }
            )
        normalized.sort(key=lambda item: item["name"].lower())
        return normalized

    # ------------------------------------------------------------------
    # Accessors
    # ------------------------------------------------------------------
    def list_categories(self) -> List[str]:
        return [category["name"] for category in self.data.get("categories", [])]

    def list_brands(self, category_name: str) -> List[str]:
        category = self._get_category(category_name)
        if not category:
            return []
        return [brand["name"] for brand in category.get("brands", [])]

    def list_models(self, category_name: str, brand_name: str) -> List[str]:
        brand = self._get_brand(category_name, brand_name)
        if not brand:
            return []
        return [model["name"] for model in brand.get("models", [])]

    def iter_model_records(self) -> Iterator[ModelRecord]:
        for category in self.data.get("categories", []):
            category_name = category.get("name")
            if not isinstance(category_name, str):
                continue
            for brand in category.get("brands", []):
                brand_name = brand.get("name")
                if not isinstance(brand_name, str):
                    continue
                for model in brand.get("models", []):
                    model_name = model.get("name")
                    if not isinstance(model_name, str):
                        continue
                    pattern = str(model.get("pattern") or "")
                    if not pattern:
                        pattern = generate_model_pattern(brand_name, model_name)
                    tags = model.get("tags") or [category_name, brand_name, model_name]
                    normalized_tags = _normalize_tags(tags)
                    yield ModelRecord(
                        category=category_name,
                        brand=brand_name,
                        name=model_name,
                        pattern=pattern,
                        tags=normalized_tags,
                    )

    # ------------------------------------------------------------------
    # Excel import/export
    # ------------------------------------------------------------------
    def export_to_excel(self, path: str | Path) -> None:
        """Export the current model list into an Excel workbook."""

        if Workbook is None:
            raise ValueError(
                "Експорт у Excel недоступний. Встановіть залежність 'openpyxl'."
            )

        destination = Path(path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "Models"
            sheet.append(_EXCEL_HEADERS)

            for record in self.iter_model_records():
                category = self._get_category(record.category) or {}
                brand = self._get_brand(record.category, record.brand) or {}
                category_tags = _normalize_tags(category.get("tags", []))
                brand_tags = _normalize_tags(brand.get("tags", []))
                row = [
                    record.category,
                    "; ".join(category_tags),
                    record.brand,
                    "; ".join(brand_tags),
                    record.name,
                    "; ".join(record.tags),
                    record.pattern,
                ]
                sheet.append(row)

            workbook.save(destination)
        finally:  # pragma: no branch - ensure closure under normal operation
            try:
                workbook.close()
            except Exception:  # pragma: no cover - defensive
                pass

    def import_from_excel(self, path: str | Path) -> None:
        """Replace current templates with the contents of an Excel workbook."""

        if load_workbook is None:
            raise ValueError(
                "Імпорт із Excel недоступний. Встановіть залежність 'openpyxl'."
            )

        source = Path(path)
        if not source.exists():
            raise FileNotFoundError(f"Файл '{source}' не знайдено")

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            try:
                sheet = workbook.active
            except AttributeError as exc:  # pragma: no cover - defensive
                raise ValueError("Файл Excel не містить активного аркуша") from exc

            rows = sheet.iter_rows(
                min_row=1, max_col=len(_EXCEL_HEADERS), values_only=True
            )
            try:
                header = next(rows)
            except StopIteration:
                self.data = {"categories": []}
                return

            normalized_header = [
                str(value).strip().lower() if value is not None else ""
                for value in header
            ]
            expected_header = [value.lower() for value in _EXCEL_HEADERS]
            if normalized_header[: len(expected_header)] != expected_header:
                raise ValueError(
                    "Невірний формат файлу Excel: перший рядок має містити заголовки "
                    + ", ".join(_EXCEL_HEADERS)
                )

            categories: Dict[str, Dict[str, object]] = {}
            seen_models: set[tuple[str, str, str]] = set()

            for index, row in enumerate(rows, start=2):
                if row is None:
                    continue
                values = list(row) + [None] * (len(_EXCEL_HEADERS) - len(row))
                category_name = str(values[0]).strip() if values[0] is not None else ""
                brand_name = str(values[2]).strip() if values[2] is not None else ""
                model_name = str(values[4]).strip() if values[4] is not None else ""

                if not category_name:
                    if all(cell is None or str(cell).strip() == "" for cell in values):
                        continue
                    raise ValueError(f"Рядок {index}: відсутня назва категорії")
                if not brand_name:
                    raise ValueError(f"Рядок {index}: відсутня назва бренду")
                if not model_name:
                    raise ValueError(f"Рядок {index}: відсутня назва моделі")

                key = (
                    category_name.lower(),
                    brand_name.lower(),
                    model_name.lower(),
                )
                if key in seen_models:
                    raise ValueError(
                        f"Рядок {index}: модель '{model_name}' для бренду '{brand_name}' "
                        f"у категорії '{category_name}' дублюється"
                    )
                seen_models.add(key)

                category_entry = categories.setdefault(
                    category_name,
                    {"name": category_name, "tags": [], "brands": {}},
                )
                category_tags = _split_tags(values[1])
                if category_tags:
                    category_entry["tags"].extend(category_tags)
                else:
                    category_entry["tags"].append(category_name)

                brands = category_entry.setdefault("brands", {})  # type: ignore[assignment]
                brand_entry = brands.setdefault(
                    brand_name,
                    {"name": brand_name, "tags": [], "models": {}},
                )
                brand_tags = _split_tags(values[3])
                if brand_tags:
                    brand_entry["tags"].extend(brand_tags)
                else:
                    brand_entry["tags"].append(brand_name)

                model_tags = _split_tags(values[5])
                if not model_tags:
                    model_tags = [category_name, brand_name, model_name]
                pattern_value = str(values[6]).strip() if values[6] is not None else ""
                model_data = {
                    "name": model_name,
                    "pattern": pattern_value or generate_model_pattern(brand_name, model_name),
                    "tags": model_tags,
                }
                models = brand_entry.setdefault("models", {})  # type: ignore[assignment]
                models[model_name] = model_data

            structured_categories: List[Dict[str, object]] = []
            for category_name, category_data in categories.items():
                brands_map: Dict[str, Dict[str, object]] = category_data.get("brands", {})  # type: ignore[assignment]
                brand_list: List[Dict[str, object]] = []
                for brand_name, brand_data in brands_map.items():
                    models_map: Dict[str, Dict[str, object]] = brand_data.get("models", {})  # type: ignore[assignment]
                    models_list = list(models_map.values())
                    brand_list.append(
                        {
                            "name": brand_name,
                            "tags": brand_data.get("tags", []),
                            "models": models_list,
                        }
                    )
                structured_categories.append(
                    {
                        "name": category_name,
                        "tags": category_data.get("tags", []),
                        "brands": brand_list,
                    }
                )

            self.data = {"categories": self._normalize_structure(structured_categories)}
        finally:
            try:
                workbook.close()
            except Exception:  # pragma: no cover - defensive
                pass

    # ------------------------------------------------------------------
    # Mutators
    # ------------------------------------------------------------------
    def add_category(self, name: str) -> None:
        cleaned = name.strip()
        if not cleaned:
            raise ValueError("Назва категорії не може бути порожньою")
        if self._get_category(cleaned):
            raise ValueError("Категорія вже існує")
        categories = self.data.setdefault("categories", [])
        categories.append({"name": cleaned, "tags": [cleaned], "brands": []})
        categories.sort(key=lambda item: item["name"].lower())

    def rename_category(self, old_name: str, new_name: str) -> None:
        category = self._get_category(old_name)
        if not category:
            raise ValueError("Категорію не знайдено")
        cleaned = new_name.strip()
        if not cleaned:
            raise ValueError("Нова назва не може бути порожньою")
        if any(
            other.get("name", "").strip().lower() == cleaned.lower()
            for other in self.data.get("categories", [])
            if other is not category
        ):
            raise ValueError("Категорія вже існує")
        category["name"] = cleaned
        category["tags"] = _normalize_tags(
            [cleaned if tag.lower() == old_name.lower() else tag for tag in category.get("tags", [])]
            or [cleaned]
        )
        for brand in category.get("brands", []):
            for model in brand.get("models", []):
                tags = model.get("tags", [])
                updated = [cleaned if t.lower() == old_name.lower() else t for t in tags]
                if cleaned.lower() not in {t.lower() for t in updated}:
                    updated.insert(0, cleaned)
                model["tags"] = _normalize_tags(updated)

    def delete_category(self, name: str) -> None:
        categories = self.data.get("categories", [])
        self.data["categories"] = [
            category for category in categories if category.get("name") != name
        ]

    def add_brand(self, category_name: str, name: str) -> None:
        category = self._get_category(category_name)
        if not category:
            raise ValueError("Спочатку оберіть категорію")
        cleaned = name.strip()
        if not cleaned:
            raise ValueError("Назва бренду не може бути порожньою")
        if self._get_brand(category_name, cleaned):
            raise ValueError("Бренд вже існує")
        brands = category.setdefault("brands", [])
        brands.append({"name": cleaned, "tags": [cleaned], "models": []})
        brands.sort(key=lambda item: item["name"].lower())

    def rename_brand(self, category_name: str, old_name: str, new_name: str) -> None:
        brand = self._get_brand(category_name, old_name)
        if not brand:
            raise ValueError("Бренд не знайдено")
        cleaned = new_name.strip()
        if not cleaned:
            raise ValueError("Нова назва не може бути порожньою")
        category = self._get_category(category_name)
        if category:
            for other in category.get("brands", []):
                if other is brand:
                    continue
                if str(other.get("name", "")).strip().lower() == cleaned.lower():
                    raise ValueError("Бренд вже існує")
        brand["name"] = cleaned
        brand["tags"] = _normalize_tags(
            [cleaned if tag.lower() == old_name.lower() else tag for tag in brand.get("tags", [])]
            or [cleaned]
        )
        for model in brand.get("models", []):
            model_name = model.get("name", "")
            model["pattern"] = generate_model_pattern(cleaned, str(model_name))
            tags = model.get("tags", [])
            updated = [cleaned if t.lower() == old_name.lower() else t for t in tags]
            if cleaned.lower() not in {t.lower() for t in updated}:
                updated.append(cleaned)
            model["tags"] = _normalize_tags(updated)

    def delete_brand(self, category_name: str, name: str) -> None:
        category = self._get_category(category_name)
        if not category:
            return
        category["brands"] = [
            brand for brand in category.get("brands", []) if brand.get("name") != name
        ]

    def add_model(self, category_name: str, brand_name: str, name: str) -> None:
        brand = self._get_brand(category_name, brand_name)
        if not brand:
            raise ValueError("Оберіть бренд, до якого додається модель")
        cleaned = name.strip()
        if not cleaned:
            raise ValueError("Назва моделі не може бути порожньою")
        for existing in brand.get("models", []):
            if str(existing.get("name")) == cleaned:
                raise ValueError("Така модель вже існує")
        pattern = generate_model_pattern(brand_name, cleaned)
        tags = _normalize_tags([category_name, brand_name, cleaned])
        brand.setdefault("models", []).append(
            {"name": cleaned, "pattern": pattern, "tags": tags}
        )
        brand["models"].sort(key=lambda item: item["name"].lower())

    def rename_model(
        self, category_name: str, brand_name: str, old_name: str, new_name: str
    ) -> None:
        model = self._get_model(category_name, brand_name, old_name)
        if not model:
            raise ValueError("Модель не знайдено")
        cleaned = new_name.strip()
        if not cleaned:
            raise ValueError("Нова назва не може бути порожньою")
        brand = self._get_brand(category_name, brand_name)
        if brand:
            for other in brand.get("models", []):
                if other is model:
                    continue
                if str(other.get("name", "")).strip().lower() == cleaned.lower():
                    raise ValueError("Така модель вже існує")
        model["name"] = cleaned
        pattern = generate_model_pattern(brand_name, cleaned)
        model["pattern"] = pattern
        tags = model.get("tags", [])
        updated = [cleaned if t.lower() == old_name.lower() else t for t in tags]
        if cleaned.lower() not in {t.lower() for t in updated}:
            updated.append(cleaned)
        model["tags"] = _normalize_tags(updated)

    def delete_model(self, category_name: str, brand_name: str, name: str) -> None:
        brand = self._get_brand(category_name, brand_name)
        if not brand:
            return
        brand["models"] = [
            model for model in brand.get("models", []) if model.get("name") != name
        ]

    # ------------------------------------------------------------------
    # Matching utilities
    # ------------------------------------------------------------------
    def find_best_model_match(self, product_name: str) -> Optional[ModelRecord]:
        """Return the model with the longest exact regex match for the product."""

        best: tuple[int, ModelRecord] | None = None
        haystack = product_name or ""
        for record in self.iter_model_records():
            regex = self._compile_pattern(record.pattern)
            if not regex:
                continue
            match = regex.search(haystack)
            if not match:
                continue
            length = len(match.group(0))
            if not best or length > best[0]:
                best = (length, record)
        return best[1] if best else None

    def _compile_pattern(self, pattern: str) -> Optional[re.Pattern[str]]:
        if pattern in self._regex_cache:
            return self._regex_cache[pattern]
        try:
            compiled = re.compile(pattern, flags=re.IGNORECASE)
        except re.error:
            self._regex_cache[pattern] = None  # type: ignore[assignment]
            return None
        self._regex_cache[pattern] = compiled
        return compiled

    # ------------------------------------------------------------------
    # Internal lookup helpers
    # ------------------------------------------------------------------
    def _get_category(self, name: str) -> Optional[Dict[str, object]]:
        for category in self.data.get("categories", []):
            if category.get("name") == name:
                return category
        return None

    def _get_brand(
        self, category_name: str, brand_name: str
    ) -> Optional[Dict[str, object]]:
        category = self._get_category(category_name)
        if not category:
            return None
        for brand in category.get("brands", []):
            if brand.get("name") == brand_name:
                return brand
        return None

    def _get_model(
        self, category_name: str, brand_name: str, model_name: str
    ) -> Optional[Dict[str, object]]:
        brand = self._get_brand(category_name, brand_name)
        if not brand:
            return None
        for model in brand.get("models", []):
            if model.get("name") == model_name:
                return model
        return None


__all__ = [
    "ModelRecord",
    "ModelTemplatesEditor",
    "generate_model_pattern",
]
